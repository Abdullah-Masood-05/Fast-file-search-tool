from __future__ import annotations

import hashlib
import json
import os
import sys
import threading
import time
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Dict, Iterator, List, Optional, Set, Tuple

from .config import AppConfig
from .persistent_index import PersistentIndex, MAX_STORED_CONTENT_CHARS
from .utils.file_utils import ExcludeMatcher, build_file_info, is_binary_file

# Well-known OS folders that are never worth indexing (matched as exact names).
_SYSTEM_FOLDERS: Set[str] = {
    "$Recycle.Bin", "System Volume Information", "Windows",
    "WinSxS", "Program Files", "Program Files (x86)",
    ".Trash", ".Trashes", ".TemporaryItems",
}
# Pseudo filesystems: only skipped directly below the filesystem root on
# POSIX ("/proc", "/sys", "/dev"), never on Windows or deeper in the tree.
_POSIX_PSEUDO_ROOT_DIRS: Set[str] = {"proc", "sys", "dev"}

_IS_WINDOWS = sys.platform == "win32"
_SYSTEM_FOLDERS_FOLDED = {n.casefold() for n in _SYSTEM_FOLDERS}

_MAX_CONTENT_SIZE = 10 * 1024 * 1024  # 10 MB
_WRITE_BATCH = 2000
_CONTENT_CHUNK = 128
_PROGRESS_INTERVAL = 0.1  # seconds


@dataclass
class IndexStats:
    total_files: int = 0
    total_folders: int = 0
    total_size: int = 0
    content_indexed: int = 0
    content_skipped: int = 0
    extension_dist: Counter = field(default_factory=Counter)
    start_time: float = 0.0
    duration: float = 0.0
    # Incremental-update details (all entries seen are counted in total_*).
    added: int = 0
    updated: int = 0
    unchanged: int = 0
    removed: int = 0
    errors: int = 0
    cancelled: bool = False

    @property
    def avg_file_size(self) -> float:
        count = self.total_files
        return self.total_size / max(1, count)

    def merge(self, other: "IndexStats") -> None:
        for name in ("total_files", "total_folders", "total_size", "content_indexed",
                     "content_skipped", "added", "updated", "unchanged", "removed",
                     "errors", "duration"):
            setattr(self, name, getattr(self, name) + getattr(other, name))
        self.extension_dist.update(other.extension_dist)
        self.cancelled = self.cancelled or other.cancelled


class Indexer:
    """Walks folders and keeps a :class:`PersistentIndex` in sync with the disk.

    ``incremental_update`` (the normal path, also used at startup) only touches
    entries whose size/mtime differ from what is stored in the database, so a
    restart reuses the on-disk index. ``index_path`` forces a full re-read of a
    root. Neither ever clears other roots' entries.
    """

    def __init__(
        self,
        config: Optional[AppConfig] = None,
        persistent_index: Optional[PersistentIndex] = None,
    ):
        self.config = config or AppConfig()
        self.persistent_index = persistent_index or PersistentIndex()
        self._stats = IndexStats()
        self._exclude_patterns: List[str] = list(self.config.exclude_patterns)
        self._excludes = ExcludeMatcher(self._exclude_patterns)
        self.index_content: bool = bool(getattr(self.config, "index_content", False))
        self._progress_callback: Optional[Callable[[int, int], None]] = None
        self._done_callback: Optional[Callable[[IndexStats], None]] = None
        self._stop_event = threading.Event()
        self._running = False
        self._last_progress = 0.0

    # ------------------------------------------------------------------ control

    def on_progress(self, cb: Callable[[int, int], None]) -> None:
        self._progress_callback = cb

    def on_done(self, cb: Callable[[IndexStats], None]) -> None:
        self._done_callback = cb

    def stop(self) -> None:
        """Request cancellation. Sticky until :meth:`reset` is called, so a stop
        that races with the start of a scan is never lost."""
        self._stop_event.set()

    def reset(self) -> None:
        self._stop_event.clear()

    @property
    def _is_running(self) -> bool:
        return self._running and not self._stop_event.is_set()

    @property
    def stopped(self) -> bool:
        return self._stop_event.is_set()

    # ------------------------------------------------------------------ filtering

    def _name_skipped(self, name: str, parent: str) -> bool:
        if name.casefold() in _SYSTEM_FOLDERS_FOLDED:
            return True
        if not _IS_WINDOWS and name in _POSIX_PSEUDO_ROOT_DIRS and parent == "/":
            return True
        return self._excludes.matches_name(name)

    def _should_skip(self, path: Path) -> bool:
        """True if ``path`` itself (by name) or any of its path components is
        excluded. Matching is by whole path component, never by substring."""
        if self.persistent_index.is_own_file(path):
            return True
        parent = str(path.parent)
        if self._name_skipped(path.name, parent):
            return True
        return self._excludes.matches_path(path)

    # ------------------------------------------------------------------ content

    def _extract_content_text(self, path: Path) -> Optional[str]:
        try:
            if not path.is_file():
                return None
            size = path.stat().st_size
        except OSError:
            return None
        if size > _MAX_CONTENT_SIZE:
            return None
        ext = path.suffix.lower()

        try:
            if ext in {".txt", ".md", ".py", ".js", ".ts", ".java", ".c", ".cpp",
                        ".h", ".hpp", ".rs", ".go", ".rb", ".php", ".css", ".html",
                        ".xml", ".json", ".yaml", ".yml", ".toml", ".cfg", ".ini",
                        ".log", ".csv", ".sh", ".bat", ".ps1", ".sql", ".r",
                        ".swift", ".kt", ".scala", ".clj", ".lua", ".pl", ".pm"}:
                try:
                    return path.read_text(encoding="utf-8", errors="replace")
                except Exception:
                    try:
                        return path.read_text(encoding="latin-1", errors="replace")
                    except Exception:
                        return None

            if ext == ".pdf":
                return self._extract_pdf(path)
            if ext in {".doc", ".docx"}:
                return self._extract_docx(path)
            if ext in {".xls", ".xlsx"}:
                return self._extract_xlsx(path)
            if ext in {".zip", ".tar", ".gz", ".bz2", ".7z", ".rar"}:
                return self._extract_archive(path)

            if ext in {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".webp"}:
                return self._extract_image_ocr(path)
        except Exception:
            pass
        return None

    def _extract_pdf(self, path: Path) -> Optional[str]:
        try:
            import pdfplumber
            text_parts = []
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages[:50]:
                    t = page.extract_text()
                    if t:
                        text_parts.append(t)
            return "\n".join(text_parts) if text_parts else None
        except ImportError:
            pass
        try:
            import PyPDF2
            text_parts = []
            with open(path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages[:50]:
                    t = page.extract_text()
                    if t:
                        text_parts.append(t)
            return "\n".join(text_parts) if text_parts else None
        except ImportError:
            return None

    def _extract_docx(self, path: Path) -> Optional[str]:
        try:
            import docx
            doc = docx.Document(path)
            return "\n".join(p.text for p in doc.paragraphs)
        except ImportError:
            return None

    def _extract_xlsx(self, path: Path) -> Optional[str]:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            text_parts = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    parts = [str(c) for c in row if c is not None]
                    if parts:
                        text_parts.append(" ".join(parts))
            wb.close()
            return "\n".join(text_parts) if text_parts else None
        except ImportError:
            return None

    def _extract_archive(self, path: Path) -> Optional[str]:
        try:
            import zipfile
            if path.suffix.lower() == ".zip":
                with zipfile.ZipFile(path, "r") as zf:
                    return "\n".join(zf.namelist())
        except Exception:
            pass
        try:
            import tarfile
            if path.suffix.lower() in {".tar", ".gz", ".bz2"}:
                with tarfile.open(path, "r") as tf:
                    return "\n".join(tf.getnames())
        except Exception:
            pass
        return None

    def _extract_image_ocr(self, path: Path) -> Optional[str]:
        try:
            import pytesseract
            from PIL import Image
            img = Image.open(path)
            text = pytesseract.image_to_string(img)
            return text.strip() or None
        except ImportError:
            return None

    def _detect_language(self, path: Path) -> Optional[str]:
        ext_map = {
            ".py": "Python", ".js": "JavaScript", ".ts": "TypeScript",
            ".java": "Java", ".c": "C", ".cpp": "C++", ".cs": "C#",
            ".go": "Go", ".rs": "Rust", ".rb": "Ruby", ".php": "PHP",
            ".swift": "Swift", ".kt": "Kotlin", ".scala": "Scala",
            ".r": "R", ".pl": "Perl", ".lua": "Lua", ".hs": "Haskell",
            ".clj": "Clojure", ".erl": "Erlang", ".ex": "Elixir",
        }
        return ext_map.get(path.suffix.lower())

    # ------------------------------------------------------------------ per-entry work

    def _index_file(
        self,
        path: Path,
        parent_path: str,
        st: Optional[os.stat_result] = None,
        is_dir: Optional[bool] = None,
    ) -> dict:
        """Build the DB record for one entry. Content is only read when
        ``index_content`` is enabled (and is capped when stored)."""
        if st is None:
            st = path.stat()
        if is_dir is None:
            is_dir = path.is_dir()
        info = build_file_info(path, st, is_dir)
        info["parent_path"] = parent_path
        info["content_hash"] = None
        info["content_text"] = None
        if self.index_content and not is_dir and not self.stopped and not is_binary_file(path):
            text = self._extract_content_text(path)
            if text is not None:
                info["content_hash"] = hashlib.sha256(
                    text.encode("utf-8", errors="replace")
                ).hexdigest()
                info["content_text"] = text[:MAX_STORED_CONTENT_CHARS]
        return info

    # ------------------------------------------------------------------ walking

    def _walk(
        self, root: str, failed_dirs: List[str]
    ) -> Iterator[Tuple[str, bool, Optional[os.stat_result]]]:
        """Yield ``(path, is_dir, stat_or_None)`` for every non-excluded entry
        below ``root`` without following symlinks/junctions.

        An unreadable directory is recorded in ``failed_dirs`` and skipped; it
        never aborts the rest of the walk. ``stat_or_None`` is None when the
        entry exists but could not be stat'ed.
        """
        stack = [root]
        has_junction = hasattr(os.DirEntry, "is_junction")
        while stack:
            if self.stopped:
                return
            current = stack.pop()
            try:
                with os.scandir(current) as it:
                    entries = list(it)
            except OSError:
                failed_dirs.append(current)
                self._stats.errors += 1
                continue
            for entry in entries:
                if self.stopped:
                    return
                name = entry.name
                if self._name_skipped(name, current):
                    continue
                path = entry.path
                if self.persistent_index.is_own_file(path):
                    continue
                if self._excludes.sequences and self._excludes.matches_path(path):
                    continue
                try:
                    is_dir = entry.is_dir(follow_symlinks=False)
                    st = entry.stat(follow_symlinks=False)
                    descend = is_dir
                    if descend and has_junction and entry.is_junction():
                        descend = False  # directory junction: avoid loops
                except OSError:
                    self._stats.errors += 1
                    yield path, False, None
                    continue
                yield path, is_dir, st
                if descend:
                    stack.append(path)

    # ------------------------------------------------------------------ scanning

    def _emit_progress(self, processed: int, total: int, force: bool = False) -> None:
        cb = self._progress_callback
        if cb is None:
            return
        now = time.monotonic()
        if not force and now - self._last_progress < _PROGRESS_INTERVAL:
            return
        self._last_progress = now
        try:
            cb(processed, max(total, processed) if total else 0)
        except Exception:
            pass

    def _build_infos(
        self, batch: List[Tuple[str, bool, os.stat_result]], pool: Optional[ThreadPoolExecutor]
    ) -> List[dict]:
        def build(item):
            path, is_dir, st = item
            if self.stopped:
                return None
            try:
                p = Path(path)
                return self._index_file(p, str(p.parent), st, is_dir)
            except Exception:
                return None

        if pool is not None and self.index_content:
            results = list(pool.map(build, batch))
        else:
            results = [build(item) for item in batch]
        return results

    def _scan(self, path: str, force: bool) -> IndexStats:
        root = Path(path)
        if not root.exists():
            raise FileNotFoundError(f"Path not found: {path}")
        root_str = str(root)

        self._stats = stats = IndexStats(start_time=time.time())
        self._running = True
        self._last_progress = 0.0
        index = self.persistent_index
        try:
            # What the database already knows about this root: path -> (size, mtime, is_folder).
            # Entries are popped as the walk meets them; what is left was deleted on disk.
            known: Dict[str, Tuple[int, int, bool]] = {
                p: (size, mtime, is_folder)
                for p, size, mtime, is_folder in index.iter_file_states(root_str)
            }
            total_estimate = len(known)
            failed_dirs: List[str] = []
            pending: List[Tuple[str, bool, os.stat_result]] = []
            processed = 0

            pool = (
                ThreadPoolExecutor(max_workers=min(8, os.cpu_count() or 4))
                if self.index_content
                else None
            )

            def flush() -> None:
                if not pending:
                    return
                infos = [i for i in self._build_infos(pending, pool) if i is not None]
                stats.errors += len(pending) - len(infos) if not self.stopped else 0
                pending.clear()
                if infos and not self.stopped:
                    for info in infos:
                        if info["content_text"] is not None:
                            stats.content_indexed += 1
                        else:
                            stats.content_skipped += 1
                    index.add_files_batch(infos)

            try:
                for entry_path, is_dir, st in self._walk(root_str, failed_dirs):
                    old = known.pop(entry_path, None)
                    processed += 1
                    if st is None:
                        # exists but unreadable: keep whatever we had
                        self._emit_progress(processed, total_estimate)
                        continue
                    if is_dir:
                        stats.total_folders += 1
                    else:
                        stats.total_files += 1
                        stats.total_size += st.st_size
                        stats.extension_dist[
                            Path(entry_path).suffix.lower() or "File"
                        ] += 1

                    unchanged = (
                        not force
                        and old is not None
                        and old[2] == is_dir
                        and (is_dir or (old[0] == st.st_size and old[1] == int(st.st_mtime)))
                    )
                    if unchanged:
                        stats.unchanged += 1
                    else:
                        if old is None:
                            stats.added += 1
                        else:
                            stats.updated += 1
                        pending.append((entry_path, is_dir, st))
                        limit = _CONTENT_CHUNK if self.index_content else _WRITE_BATCH
                        if len(pending) >= limit:
                            flush()
                    self._emit_progress(processed, total_estimate)
                flush()
            finally:
                if pool is not None:
                    pool.shutdown(wait=True, cancel_futures=True)

            if self.stopped:
                stats.cancelled = True
            else:
                # Anything still in `known` vanished from disk, unless it lives under
                # a directory we could not read (then we simply don't know).
                gone = [p for p in known if not self._under_any(p, failed_dirs)]
                if gone:
                    index.remove_files_batch(gone, include_descendants=False)
                    index.cleanup_orphan_trigrams()
                    stats.removed = len(gone)

            stats.duration = time.time() - stats.start_time
            if not stats.cancelled:
                self._record_metadata(root_str, stats)
            self._emit_progress(processed, total_estimate, force=True)
        finally:
            self._running = False

        if self._done_callback:
            try:
                self._done_callback(stats)
            except Exception:
                pass
        return stats

    @staticmethod
    def _under_any(path: str, dirs: List[str]) -> bool:
        for d in dirs:
            base = d.rstrip("/\\")
            if path == d or path.startswith(base + "/") or path.startswith(base + "\\"):
                return True
        return False

    def _record_metadata(self, root: str, stats: IndexStats) -> None:
        index = self.persistent_index
        index.update_metadata("last_updated", str(int(time.time())))
        index.update_metadata(
            "index_stats",
            json.dumps({
                "root": root,
                "total_files": stats.total_files,
                "total_folders": stats.total_folders,
                "total_size": stats.total_size,
                "added": stats.added,
                "updated": stats.updated,
                "removed": stats.removed,
                "unchanged": stats.unchanged,
                "errors": stats.errors,
                "duration": stats.duration,
            }),
        )

    def index_path(self, path: str) -> IndexStats:
        """Full (re)index of one root: every entry is re-read and rewritten and
        stale entries under the root are removed. Other roots are untouched."""
        return self._scan(path, force=True)

    def incremental_update(self, path: str) -> IndexStats:
        """Bring the stored index of one root up to date, touching only entries
        that are new, changed (size/mtime differ from the DB) or deleted."""
        return self._scan(path, force=False)

    @property
    def stats(self) -> IndexStats:
        return self._stats

    def get_index_stats(self) -> dict:
        return self.persistent_index.get_index_stats()
